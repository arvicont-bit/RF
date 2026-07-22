"""Seed chart_accounts din 'Plan de Conturi - Mapare Cod Raportare.xlsx' (499 conturi)."""
import openpyxl
import os
import sys

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MASTER_FILE = os.path.join(PROJECT_ROOT, 'data', 'seed', 'Plan de Conturi - Mapare Cod Raportare.xlsx')

sys.path.insert(0, PROJECT_ROOT)
from app.db import get_conn, init_db  # noqa: E402


def seed():
    init_db()
    conn = get_conn()
    cur = conn.cursor()
    cur.execute('SELECT COUNT(*) FROM chart_accounts')
    if cur.fetchone()[0] > 0:
        print('chart_accounts deja populat, sar peste seed.')
        conn.close()
        return

    wb = openpyxl.load_workbook(MASTER_FILE, data_only=True)
    ws = wb['Plan conturi - Mapare']
    rows = []
    for r in range(5, ws.max_row + 1):
        cont = ws.cell(row=r, column=2).value
        if cont is None:
            continue
        rows.append((
            str(ws.cell(row=r, column=1).value or ''),   # clasa
            str(cont).strip(),                            # cont
            str(ws.cell(row=r, column=3).value or ''),   # grad
            ws.cell(row=r, column=4).value,               # denumire
            ws.cell(row=r, column=5).value,               # tip
            ws.cell(row=r, column=6).value,               # cod_mapare
            ws.cell(row=r, column=7).value,               # descriere_cod
            ws.cell(row=r, column=8).value,               # categorie
            ws.cell(row=r, column=9).value,               # subcategorie
            ws.cell(row=r, column=10).value,              # tratament_special
            ws.cell(row=r, column=11).value,              # sursa
        ))

    cur.executemany(
        '''INSERT INTO chart_accounts
           (clasa, cont, grad, denumire, tip, cod_mapare, descriere_cod, categorie, subcategorie, tratament_special, sursa)
           VALUES (?,?,?,?,?,?,?,?,?,?,?)''',
        rows,
    )
    conn.commit()
    print(f'Seed complet: {len(rows)} conturi inserate in chart_accounts.')
    conn.close()


if __name__ == '__main__':
    seed()
